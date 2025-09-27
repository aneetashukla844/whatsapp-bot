from flask import Flask, request
import json
import os
import requests
from openpyxl import Workbook, load_workbook

app = Flask(__name__)

DATA_FILE = "appointments.json"
EXCEL_FILE = "appointments.xlsx"

# Load appointments from JSON
if os.path.exists(DATA_FILE):
    with open(DATA_FILE, "r") as f:
        appointments = json.load(f)
else:
    appointments = {}

# Messages for multilingual support
messages = {
    "ask_name": {"en": "Enter your full name:", "hi": "कृपया अपना पूरा नाम दर्ज करें:"},
    "ask_mobile": {"en": "Enter your mobile number:", "hi": "कृपया अपना मोबाइल नंबर दर्ज करें:"},
    "ask_age": {"en": "Enter your age:", "hi": "कृपया अपनी आयु दर्ज करें:"},
    "ask_address": {"en": "Enter your address:", "hi": "कृपया अपना पता दर्ज करें:"},
    "already_booked": {"en": "You already have a token:", "hi": "आपके पास पहले से ही एक टोकन है:"},
    "appointment_booked": {"en": "Appointment booked! Your token number is", "hi": "आपकी अपॉइंटमेंट बुक हो गई है! आपका टोकन नंबर है"}
}

VERIFY_TOKEN = "clinic-bot"
ACCESS_TOKEN = "YOUR_META_ACCESS_TOKEN"
PHONE_NUMBER_ID = "YOUR_PHONE_NUMBER_ID"

# Create Excel file if it does not exist
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.append(["Token", "Name", "Mobile", "Age", "Address"])
    wb.save(EXCEL_FILE)

# Detect language (simple check for Hindi)
def detect_language(text):
    for c in text:
        if '\u0900' <= c <= '\u097F':
            return "hi"
    return "en"

# Send message using Meta API
def send_whatsapp_message(to, message):
    url = f"https://graph.facebook.com/v17.0/{PHONE_NUMBER_ID}/messages"
    headers = {
        "Authorization": f"Bearer {ACCESS_TOKEN}",
        "Content-Type": "application/json"
    }
    payload = {
        "messaging_product": "whatsapp",
        "to": to,
        "type": "text",
        "text": {"body": message}
    }
    requests.post(url, headers=headers, json=payload)

# Send multilingual message
def send_multilingual_message(to, key, lang="en", extra=""):
    msg = messages[key][lang]
    if extra:
        msg += f" {extra}"
    send_whatsapp_message(to, msg)

# Save JSON data
def save_appointments():
    with open(DATA_FILE, "w") as f:
        json.dump(appointments, f, indent=2)

# Save appointment to Excel and generate token
def save_to_excel(user_data):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    token_number = len(ws['A'])  # Next row as token
    ws.append([token_number, user_data["name"], user_data["mobile"], user_data["age"], user_data["address"]])
    wb.save(EXCEL_FILE)
    return token_number

# Webhook verification
@app.route('/webhook', methods=['GET'])
def verify():
    token = request.args.get("hub.verify_token")
    challenge = request.args.get("hub.challenge")
    if token == VERIFY_TOKEN:
        return challenge
    return "Invalid verification token", 403

# Webhook to receive messages
@app.route('/webhook', methods=['POST'])
def webhook():
    data = request.get_json()
    try:
        entry = data['entry'][0]
        changes = entry['changes'][0]
        value = changes['value']
        messages_list = value.get('messages')

        if messages_list:
            for message in messages_list:
                phone_number = message['from']
                text = message['text']['body'].strip()
                lang = detect_language(text)

                if phone_number not in appointments:
                    appointments[phone_number] = {"step": "ask_name"}
                    save_appointments()
                    send_multilingual_message(phone_number, "ask_name", lang)

                else:
                    user_data = appointments[phone_number]
                    step = user_data.get("step")

                    if step == "ask_name":
                        user_data["name"] = text
                        user_data["step"] = "ask_mobile"
                        save_appointments()
                        send_multilingual_message(phone_number, "ask_mobile", lang)

                    elif step == "ask_mobile":
                        user_data["mobile"] = text
                        user_data["step"] = "ask_age"
                        save_appointments()
                        send_multilingual_message(phone_number, "ask_age", lang)

                    elif step == "ask_age":
                        user_data["age"] = text
                        user_data["step"] = "ask_address"
                        save_appointments()
                        send_multilingual_message(phone_number, "ask_address", lang)

                    elif step == "ask_address":
                        user_data["address"] = text
                        token_number = save_to_excel(user_data)
                        user_data["token"] = token_number
                        user_data["step"] = "booked"
                        save_appointments()
                        send_multilingual_message(phone_number, "appointment_booked", lang, token_number)

                    elif step == "booked":
                        send_multilingual_message(phone_number, "already_booked", lang, user_data["token"])

    except Exception as e:
        print("Error:", e)

    return "EVENT_RECEIVED", 200

if __name__ == '__main__':
    app.run(debug=True)
